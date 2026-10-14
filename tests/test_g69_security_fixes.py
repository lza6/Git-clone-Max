"""G69 安全/正确性修复守卫（独立审计线程报出的 M-1/M-2/M-3/L-1）。

每条都先「实测过问题」，再修，再固化为守卫：
- G69-8（M-1）本地 HTTP API 不校验 Host/Origin → DNS rebinding 可读数据
- G69-9（M-2）MCP `export_report` 的 out_dir 可写任意绝对路径
- G69-7（M-3）报表公式注入（owner 来自用户粘贴的清单）
- G69-10（L-1）`/api/report?format=csv` 的列取自 sync_history、行取自 repos → 恒空
"""
from __future__ import annotations

import csv
import io
import json
import os
import pathlib
import socket
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


def _free_port() -> int:
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


class _ServerCase(unittest.TestCase):
    """起一个真实的本地 API 服务（回环绑定），用于端到端校验。"""

    @classmethod
    def setUpClass(cls):
        from gcm.db.repo_db import Database
        from gcm.http_api import create_server
        from gcm.models import RepoSpec, SyncAction, SyncStatus

        cls.tmp = pathlib.Path(tempfile.mkdtemp(prefix="gcm_g69api_"))
        db = Database(cls.tmp / "repos.db")
        spec = RepoSpec(owner="a", repo="b", url_https="https://github.com/a/b.git",
                        folder_name="a__b")
        db.upsert_repo(spec, str(cls.tmp / "clones" / "a__b"))
        db.add_sync_history(1, SyncStatus.SUCCESS, SyncAction.CLONED,
                            message="完成", commits=3)
        db.close()

        cls.port = _free_port()
        cls.srv = create_server(f"127.0.0.1:{cls.port}", cls.tmp, allow_mutate=False)
        cls.thread = threading.Thread(target=cls.srv.serve_forever, daemon=True)
        cls.thread.start()
        time.sleep(0.4)
        cls.base = f"http://127.0.0.1:{cls.port}"

    @classmethod
    def tearDownClass(cls):
        try:
            cls.srv.shutdown()
        except Exception:
            pass

    def _req(self, path: str, headers: dict | None = None):
        req = urllib.request.Request(self.base + path, headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=10) as r:
                return r.status, r.read().decode("utf-8", "replace")
        except urllib.error.HTTPError as e:
            return e.code, e.read().decode("utf-8", "replace")


class TestHostOriginHardening(_ServerCase):
    """G69-8：非回环 Host / 跨站 Origin 必须被拒。"""

    def test_forged_host_is_rejected(self):
        """DNS rebinding 场景：域名解析到 127.0.0.1 后 Host 仍是攻击者域名。"""
        status, body = self._req("/api/repos", {"Host": "evil.example.com"})
        self.assertEqual(status, 403, f"伪造 Host 未被拒绝：{status} {body[:120]}")
        self.assertIn("rebinding", body)

    def test_cross_site_origin_is_rejected(self):
        status, body = self._req("/api/repos", {"Origin": "https://evil.example.com"})
        self.assertEqual(status, 403, f"跨站 Origin 未被拒绝：{status} {body[:120]}")

    def test_loopback_requests_still_work(self):
        """反向约束：正常回环请求不得被误伤。"""
        self.assertEqual(self._req("/api/repos")[0], 200)
        self.assertEqual(
            self._req("/api/repos", {"Origin": f"http://127.0.0.1:{self.port}"})[0], 200)
        self.assertEqual(self._req("/api/repos", {"Host": f"localhost:{self.port}"})[0], 200)

    def test_missing_host_is_rejected(self):
        """缺 Host 头 → 保守拒绝（原始 socket 构造，绕过 urllib 自动加 Host）。"""
        raw = (f"GET /api/repos HTTP/1.1\r\n"          # 故意不带 Host
               f"Connection: close\r\n\r\n").encode()
        s = socket.create_connection(("127.0.0.1", self.port), timeout=10)
        try:
            s.sendall(raw)
            data = s.recv(200).decode("utf-8", "replace")
        finally:
            s.close()
        self.assertIn("403", data.splitlines()[0] if data else "", f"缺 Host 未拒绝：{data[:80]}")


class TestReportHasRealColumns(_ServerCase):
    """G69-10：报表接口必须返回真实列（原来除 owner/repo/host 外全空）。"""

    def test_csv_contains_history_fields(self):
        status, body = self._req("/api/report?format=csv")
        self.assertEqual(status, 200)
        text = body.lstrip("﻿")
        rows = list(csv.reader(io.StringIO(text)))
        self.assertEqual(rows[0][:3], ["owner", "repo", "host"])
        self.assertIn("status", rows[0])
        self.assertGreaterEqual(len(rows), 2, "应有数据行")
        data = rows[1]
        self.assertEqual(data[:3], ["a", "b", "github.com"])
        self.assertEqual(data[rows[0].index("status")], "success",
                         "status 列恒空 —— 说明又用了 repos 表取 history 字段")
        self.assertEqual(data[rows[0].index("commits")], "3")


class TestMcpOutDirIsConfined(unittest.TestCase):
    """G69-9：MCP 导出目录必须限定在数据目录内。"""

    def setUp(self):
        from gcm.db.repo_db import Database

        self.tmp = pathlib.Path(tempfile.mkdtemp(prefix="gcm_g699_"))
        self.db = Database(self.tmp / "repos.db")

    def tearDown(self):
        try:
            self.db.close()
        except Exception:
            pass

    def _export(self, out_dir):
        from gcm.mcp_server import _export_report

        return _export_report(self.db, "csv", out_dir)

    def test_outside_absolute_path_is_rejected(self):
        outside = pathlib.Path(tempfile.mkdtemp()) / "evil"
        with self.assertRaises(ValueError):
            self._export(str(outside))
        self.assertFalse(outside.exists(), "越界目录竟被创建")

    def test_parent_traversal_is_rejected(self):
        with self.assertRaises(ValueError):
            self._export(str(self.tmp / ".." / "escape"))

    def test_relative_and_default_are_allowed(self):
        msg = self._export("sub/dir")
        self.assertIn("sub", msg)
        self.assertTrue((self.tmp / "sub" / "dir" / "repos_report.csv").exists())
        self._export(None)   # 默认 reports/
        self.assertTrue((self.tmp / "reports" / "repos_report.csv").exists())


class TestSpreadsheetFormulaInjection(unittest.TestCase):
    """G69-7：以 = + - @ 开头的单元格内容必须被转义成纯文本。"""

    def setUp(self):
        from gcm.db.repo_db import Database
        from gcm.models import RepoSpec, SyncAction, SyncStatus

        self.tmp = pathlib.Path(tempfile.mkdtemp(prefix="gcm_g697_"))
        self.db = Database(self.tmp / "repos.db")
        spec = RepoSpec(owner="=HYPERLINK(A1)", repo="b",
                        url_https="https://github.com/x/b.git", folder_name="x__b")
        self.db.upsert_repo(spec, str(self.tmp / "clones" / "x__b"))
        self.db.add_sync_history(1, SyncStatus.SUCCESS, SyncAction.CLONED,
                                 message="+cmd|'/c calc'!A0", commits=1)

    def tearDown(self):
        try:
            self.db.close()
        except Exception:
            pass

    def test_sanitize_cell_covers_all_prefixes(self):
        from gcm.reports import sanitize_cell

        for v in ("=HYPERLINK(A1)", "+1+1", "-2+2", "@SUM(A1)", "\t=x", "\r=x"):
            self.assertTrue(sanitize_cell(v).startswith("'"), f"{v!r} 未被转义")
        for v in ("normal", "中文", "", "a=b"):
            self.assertEqual(sanitize_cell(v), v, f"{v!r} 被误伤")

    def test_csv_escapes_formula_prefix(self):
        from gcm.reports import export_csv

        out = self.tmp / "r.csv"
        export_csv(self.db, out)
        text = out.read_text(encoding="utf-8-sig")
        self.assertIn("'=HYPERLINK(A1)", text, "CSV 未转义公式前缀")
        self.assertNotIn(",=HYPERLINK", text, "公式前缀仍裸奔")

    def test_xlsx_escapes_formula_prefix(self):
        try:
            import openpyxl  # noqa: F401
        except Exception:
            self.skipTest("未安装 openpyxl")
        from gcm.reports import export_xlsx

        out = self.tmp / "r.xlsx"
        export_xlsx(self.db, out)
        import openpyxl
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        col = headers.index("owner") + 1
        cell = ws.cell(row=2, column=col)
        self.assertEqual(cell.data_type, "s",
                         f"xlsx 单元格仍是公式（data_type={cell.data_type}）")
        self.assertTrue(str(cell.value).startswith("'"))


class TestCredentialsStrippedAtIngest(unittest.TestCase):
    """G69-11：本地扫描带出来的内嵌凭据不得进入 DB / 导出 / 剪贴板。

    实测原状：remote `https://ghp_SECRET@github.com/team/private.git` →
    ① 扫描兜底正则把 `user:pass@host` 整段当 host，url_https 里带着 token
    ② DB `url` 列原样存 token  ③ 元数据导出 JSON（定位「便于分享」）带着 token
    ④ 详情页「复制克隆命令」把它写进剪贴板。拿到即可读写该用户私有仓库。
    """

    def setUp(self):
        from gcm.db.repo_db import Database

        self.tmp = pathlib.Path(tempfile.mkdtemp(prefix="gcm_g6911_"))
        self.db = Database(self.tmp / "repos.db")

    def tearDown(self):
        try:
            self.db.close()
        except Exception:
            pass

    def test_scanner_fallback_does_not_treat_userinfo_as_host(self):
        from gcm.app.scanner import parse_remote_url

        owner, repo, url, host = parse_remote_url(
            "https://ghp_SECRET_PAT_ABC123@github.com/team/private.git")
        self.assertNotIn("ghp_", url, f"url_https 仍含 token：{url}")
        self.assertNotIn("ghp_", host, f"host 被 userinfo 污染：{host}")
        self.assertEqual(host, "github.com")
        self.assertEqual((owner, repo), ("team", "private"))

    def test_scanner_keeps_normal_urls_intact(self):
        """反向约束：正常 URL 的解析结果不得被我的正则改动破坏。"""
        from gcm.app.scanner import parse_remote_url

        for raw, exp in (
            ("https://github.com/o/r.git", ("o", "r", "github.com")),
            ("git@github.com:o/r.git", ("o", "r", "github.com")),
            ("ssh://git@gitlab.company.com/grp/proj.git", ("grp", "proj", "gitlab.company.com")),
        ):
            o, r, url, h = parse_remote_url(raw)
            self.assertEqual((o, r, h), exp, f"{raw} 解析被破坏")
            self.assertTrue(url.startswith("https://"), url)

    def test_db_url_column_has_no_credentials(self):
        from gcm.models import RepoSpec

        self.db.upsert_repo(RepoSpec(owner="a", repo="b",
                                     url_https="https://ghp_SECRET@github.com/a/b.git",
                                     folder_name="a__b"), str(self.tmp / "p"))
        row = self.db._conn.execute("SELECT url FROM repos").fetchone()
        self.assertNotIn("ghp_", row["url"], f"DB 里仍有 token：{row['url']}")
        self.assertEqual(row["url"], "https://github.com/a/b.git")

    def test_bulk_upsert_also_strips(self):
        from gcm.models import RepoSpec

        self.db.bulk_upsert_repos([
            (RepoSpec(owner="c", repo="dd",
                      url_https="https://glpat-TOK@gitlab.com/c/dd.git",
                      folder_name="c__dd"), str(self.tmp / "q"), "gitlab.com", None)])
        row = self.db._conn.execute("SELECT url FROM repos").fetchone()
        self.assertNotIn("glpat-", row["url"], f"批量入库仍有 token：{row['url']}")

    def test_metadata_export_has_no_credentials(self):
        from gcm.models import RepoSpec

        self.db.upsert_repo(RepoSpec(owner="a", repo="b",
                                     url_https="https://ghp_X@github.com/a/b.git",
                                     folder_name="a__b"), str(self.tmp / "p"))
        out = self.tmp / "meta.json"
        self.db.export_metadata_json(out)
        text = out.read_text(encoding="utf-8")
        self.assertNotIn("ghp_X", text, "可分享的元数据文件里仍有 token")


class TestTokenEncryptionIsReal(unittest.TestCase):
    """G69-12：token 落盘必须是**真加密**，不能是「源码里写着密钥的异或」。

    实测原状（打包产物路径）：依赖只有 PyQt6、spec hiddenimports 为空
    → win32crypt/keyring 都不可用 → 后端恒为 `xorfixed`，
    用仓库源码里的 `_XOR_KEY` 即可离线还原出全部 token。
    修复：ctypes 直调 crypt32 的 DPAPI（零新依赖），后端变为 `dpapi_ctypes`。
    """

    def test_backend_is_not_weak_xor_on_windows(self):
        import os as _os

        from gcm.db import settings as S

        if _os.name != "nt":
            self.skipTest("DPAPI 是 Windows 特有")
        self.assertEqual(S._get_backend(), "dpapi_ctypes",
                         "Windows 上未用上 DPAPI —— token 又退回可逆混淆了")

    def test_source_key_cannot_recover_ciphertext(self):
        """决定性断言：用仓库里公开的 `_XOR_KEY` 解**新**密文必须失败。"""
        from gcm.db import settings as S

        if S._get_backend() != "dpapi_ctypes":
            self.skipTest("非 DPAPI 后端，本条不适用")
        import base64

        plain = "ghp_SUPERSECRET_TOKEN_1234567890"
        enc = S._encrypt_token(plain)
        data = base64.b64decode(enc[len(S._ENC_PREFIX):])
        k = S._XOR_KEY
        recovered = bytes(b ^ k[i % len(k)] for i, b in enumerate(data))
        self.assertNotEqual(recovered, plain.encode("utf-8"),
                            "用源码里的公开密钥仍能还原 token —— 等于没加密")

    def test_round_trip_and_no_plaintext_on_disk(self):
        import json

        from gcm.db.settings import SettingsStore

        d = pathlib.Path(tempfile.mkdtemp(prefix="gcm_g6912_"))
        store = SettingsStore(d / "settings.json")
        plain = "ghp_SUPERSECRET_TOKEN_1234567890"
        st = store.load()
        st.token = plain
        st.host_tokens = {"gitlab.com": "glpat-PRIVATE"}
        store.save(st)

        text = (d / "settings.json").read_text(encoding="utf-8")
        self.assertNotIn(plain, text, "token 明文落盘")
        self.assertNotIn("glpat-PRIVATE", text, "host_tokens 明文落盘")
        raw = json.loads(text)
        self.assertTrue(str(raw["token"]).startswith("enc:"), raw["token"][:20])
        back = store.load()
        self.assertEqual(back.token, plain, "加密后无法读回")
        self.assertEqual(back.host_tokens, {"gitlab.com": "glpat-PRIVATE"})

    def test_weak_backend_is_disclosed(self):
        """若将来又落到弱后端，必须给出用户可见的说明（不能继续假装已加密）。"""
        from gcm.db import settings as S

        warn = S.crypto_warning()
        if S._get_backend() == "xorfixed":
            self.assertTrue(warn, "落到固定密钥异或却不告知用户")
            self.assertIn("不等于加密", warn)
        else:
            self.assertEqual(warn, "", "已真加密却仍报警")

    def test_dpapi_module_round_trip(self):
        from gcm.util import dpapi

        if not dpapi.available():
            self.skipTest("本环境无 DPAPI")
        for payload in (b"x", "中文😀".encode("utf-8"), b"a" * 500):
            self.assertEqual(dpapi.unprotect(dpapi.protect(payload)), payload)


if __name__ == "__main__":
    unittest.main()
